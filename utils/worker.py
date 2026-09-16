import random
import re
import time

from openpyxl import load_workbook

try:
    from new_parser.utils.fsa_api import collect_record_ids, extract_record_fields, fetch_nsi_labels, fetch_record
    from new_parser.utils.fsa_constants import (
        DELAY_BETWEEN_REQUESTS_SEC_MAX,
        DELAY_BETWEEN_REQUESTS_SEC_MIN,
        DOC_TYPE_DECLARATION,
    )
except ModuleNotFoundError:
    from utils.fsa_api import collect_record_ids, extract_record_fields, fetch_nsi_labels, fetch_record
    from utils.fsa_constants import (
        DELAY_BETWEEN_REQUESTS_SEC_MAX,
        DELAY_BETWEEN_REQUESTS_SEC_MIN,
        DOC_TYPE_DECLARATION,
    )

# Символы, которые openpyxl/Excel не допускают в ячейках.
_ILLEGAL_XLSX = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _clean_cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        return _ILLEGAL_XLSX.sub('', value)
    return value


def _pause():
    time.sleep(random.uniform(DELAY_BETWEEN_REQUESTS_SEC_MIN, DELAY_BETWEEN_REQUESTS_SEC_MAX))


def _cancelled(host) -> bool:
    return bool(getattr(host, 'is_cancelled', lambda: False)())


def execute_all_passes(
    host,
    path_xlsx,
    token,
    start_date,
    end_date,
    tech_key,
    doc_type=DOC_TYPE_DECLARATION,
    group_ids=None,
):
    try:
        ids = collect_record_ids(
            token,
            start_date,
            end_date,
            tech_key,
            doc_type,
            group_ids=group_ids,
            on_count=lambda count, total: host.progress_q.put(('harvest', count, total)),
            should_stop=lambda: _cancelled(host),
        )
        if _cancelled(host):
            host.progress_q.put(('cancelled',))
            return
        if not ids:
            host.progress_q.put(('error', 'По заданным параметрам ID не найдены'))
            return

        host.progress_q.put(('harvest', len(ids), len(ids)))

        try:
            book = load_workbook(path_xlsx)
        except Exception as exc:
            raise RuntimeError(f'Не удалось открыть файл Excel: {exc}') from exc
        grid = book.active
        host.slot_total = len(ids)
        host.slot_done = 0
        host.progress_q.put(('parse', 0, host.slot_total))
        row_errors: list[str] = []

        for pub_id in ids:
            if _cancelled(host):
                book.save(path_xlsx)
                host.progress_q.put(('cancelled',))
                return
            try:
                if pub_id in [str(grid[f'O{k}'].value) for k in range(2, grid.max_row + 1)]:
                    host.slot_done += 1
                    host.progress_q.put(('parse', host.slot_done, host.slot_total))
                    continue

                anchor_row = None
                for row_idx in range(2, grid.max_row + 2):
                    if grid[f'C{row_idx}'].value is None:
                        anchor_row = row_idx
                        break
                if anchor_row is None:
                    anchor_row = grid.max_row + 1

                record = fetch_record(token, pub_id, doc_type)
                fields = extract_record_fields(record)
                labels = fetch_nsi_labels(
                    token,
                    scheme_ref=fields.get('scheme_ref'),
                    record_id=pub_id,
                    doc_type=doc_type,
                )

                title = fields.get('declaration_period') or fields.get('number') or pub_id
                grid[f'C{anchor_row}'] = _clean_cell(title)
                grid[f'D{anchor_row}'] = _clean_cell(labels.get('scheme'))
                grid[f'E{anchor_row}'] = _clean_cell(fields.get('status'))
                grid[f'F{anchor_row}'] = _clean_cell(fields.get('applicant'))
                grid[f'G{anchor_row}'] = _clean_cell(fields.get('inn'))
                grid[f'H{anchor_row}'] = _clean_cell(fields.get('manufacturer'))
                grid[f'I{anchor_row}'] = _clean_cell(fields.get('address'))
                grid[f'J{anchor_row}'] = _clean_cell(fields.get('product_name'))
                grid[f'K{anchor_row}'] = _clean_cell(fields.get('document'))
                grid[f'L{anchor_row}'] = _clean_cell(fields.get('standards'))
                grid[f'M{anchor_row}'] = _clean_cell(fields.get('testing_labs'))
                grid[f'N{anchor_row}'] = _clean_cell(fields.get('testing_protocols'))
                grid[f'O{anchor_row}'] = pub_id
                grid[f'P{anchor_row}'] = _clean_cell(fields.get('object_type'))

                host.slot_done += 1
                host.progress_q.put(('parse', host.slot_done, host.slot_total))
                book.save(path_xlsx)
                _pause()
            except Exception as exc:
                row_errors.append(f'ID {pub_id}: {exc}')
                host.slot_done += 1
                host.progress_q.put(('parse', host.slot_done, host.slot_total))
                continue

        book.save(path_xlsx)
        if row_errors:
            host.progress_q.put(
                ('warn', f'Готово с ошибками ({len(row_errors)}):\n' + '\n'.join(row_errors))
            )
    except Exception as exc:
        if _cancelled(host):
            host.progress_q.put(('cancelled',))
            return
        host.progress_q.put(('error', str(exc)))
